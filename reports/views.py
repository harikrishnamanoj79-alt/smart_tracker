from django.shortcuts import render
from django.views import View
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Sum
from django.http import HttpResponse
from django.template.loader import get_template
from xhtml2pdf import pisa
import openpyxl
import csv
from io import BytesIO
import datetime
from income.models import Income
from expenses.models import Expense, Category
from savings.models import SavingsGoal
from budgets.models import Budget

# Helper to render PDF
def render_to_pdf(template_src, context_dict={}):
    template = get_template(template_src)
    html = template.render(context_dict)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode("utf-8")), result)
    if not pdf.err:
        return HttpResponse(result.getvalue(), content_type='application/pdf')
    return HttpResponse("Error generating PDF", status=500)

class ReportDashboardView(LoginRequiredMixin, View):
    def get(self, request):
        user = request.user
        today = datetime.date.today()
        
        # Selected timeframes
        selected_date_str = request.GET.get('daily_date', today.strftime('%Y-%m-%d'))
        selected_month = int(request.GET.get('monthly_month', today.month))
        selected_year = int(request.GET.get('monthly_year', today.year))
        selected_yearly_year = int(request.GET.get('yearly_year', today.year))
        
        # Parse daily date
        try:
            daily_date = datetime.datetime.strptime(selected_date_str, '%Y-%m-%d').date()
        except ValueError:
            daily_date = today
            
        # Daily calculations
        daily_income = Income.objects.filter(user=user, date=daily_date).aggregate(total=Sum('amount'))['total'] or 0
        daily_expense = Expense.objects.filter(user=user, date=daily_date).aggregate(total=Sum('amount'))['total'] or 0
        daily_balance = daily_income - daily_expense
        
        # Monthly calculations
        monthly_income_sum = Income.objects.filter(user=user, date__month=selected_month, date__year=selected_year).aggregate(total=Sum('amount'))['total'] or 0
        monthly_expense_sum = Expense.objects.filter(user=user, date__month=selected_month, date__year=selected_year).aggregate(total=Sum('amount'))['total'] or 0
        
        category_expenses = []
        for cat in Category.objects.filter(type='Expense'):
            spent = Expense.objects.filter(
                user=user,
                category=cat,
                date__month=selected_month,
                date__year=selected_year
            ).aggregate(total=Sum('amount'))['total'] or 0
            if spent > 0:
                category_expenses.append({
                    'category': cat.name,
                    'amount': spent
                })
                
        # Monthly Savings
        monthly_savings = SavingsGoal.objects.filter(user=user, target_date__month=selected_month, target_date__year=selected_year)
        
        # Yearly calculations
        yearly_income = Income.objects.filter(user=user, date__year=selected_yearly_year).aggregate(total=Sum('amount'))['total'] or 0
        yearly_expense = Expense.objects.filter(user=user, date__year=selected_yearly_year).aggregate(total=Sum('amount'))['total'] or 0
        yearly_balance = yearly_income - yearly_expense
        
        # Yearly Budget Analysis
        yearly_budgets = Budget.objects.filter(user=user, year=selected_yearly_year)
        budget_analysis = []
        for b in yearly_budgets:
            spent = Expense.objects.filter(
                user=user,
                category=b.category,
                date__year=selected_yearly_year,
                date__month=b.month
            ).aggregate(total=Sum('amount'))['total'] or 0
            
            budget_analysis.append({
                'category': b.category.name,
                'month': datetime.date(2000, b.month, 1).strftime('%B'),
                'limit': b.monthly_limit,
                'spent': spent,
                'status': 'Exceeded' if spent > b.monthly_limit else 'Under Control'
            })
            
        months = [(i, datetime.date(2000, i, 1).strftime('%B')) for i in range(1, 13)]
        years = [y for y in range(today.year - 3, today.year + 4)]

        context = {
            'daily_date': daily_date.strftime('%Y-%m-%d'),
            'daily_income': daily_income,
            'daily_expense': daily_expense,
            'daily_balance': daily_balance,
            'selected_month': selected_month,
            'selected_year': selected_year,
            'monthly_income_sum': monthly_income_sum,
            'monthly_expense_sum': monthly_expense_sum,
            'category_expenses': category_expenses,
            'monthly_savings': monthly_savings,
            'selected_yearly_year': selected_yearly_year,
            'yearly_income': yearly_income,
            'yearly_expense': yearly_expense,
            'yearly_balance': yearly_balance,
            'budget_analysis': budget_analysis,
            'months': months,
            'years': years
        }
        return render(request, 'reports/reports_dashboard.html', context)


# Exports
class ExportPDFView(LoginRequiredMixin, View):
    def get(self, request):
        user = request.user
        today = datetime.date.today()
        
        # Fetch all income & expenses to build a full statement
        incomes = Income.objects.filter(user=user).select_related('category')
        expenses = Expense.objects.filter(user=user).select_related('category')
        
        total_income = incomes.aggregate(total=Sum('amount'))['total'] or 0
        total_expense = expenses.aggregate(total=Sum('amount'))['total'] or 0
        
        context = {
            'user': user,
            'date': today,
            'incomes': incomes,
            'expenses': expenses,
            'total_income': total_income,
            'total_expense': total_expense,
            'balance': total_income - total_expense
        }
        return render_to_pdf('reports/pdf_template.html', context)


class ExportExcelView(LoginRequiredMixin, View):
    def get(self, request):
        user = request.user
        wb = openpyxl.Workbook()
        
        # Sheet 1: Income
        ws1 = wb.active
        ws1.title = "Income"
        ws1.append(["Date", "Source", "Category", "Amount", "Description"])
        incomes = Income.objects.filter(user=user)
        for inc in incomes:
            ws1.append([inc.date.strftime('%Y-%m-%d'), inc.source, inc.category.name, float(inc.amount), inc.description or ""])
            
        # Sheet 2: Expenses
        ws2 = wb.create_sheet(title="Expenses")
        ws2.append(["Date", "Category", "Amount", "Description", "Payment Method"])
        expenses = Expense.objects.filter(user=user)
        for exp in expenses:
            ws2.append([exp.date.strftime('%Y-%m-%d'), exp.category.name, float(exp.amount), exp.description or "", exp.payment_method.name if exp.payment_method else ""])
            
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Financial_Report.xlsx"'
        wb.save(response)
        return response


class ExportCSVView(LoginRequiredMixin, View):
    def get(self, request):
        user = request.user
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="Transactions_Statement.csv"'
        
        writer = csv.writer(response)
        writer.writerow(["Type", "Date", "Title/Source", "Category", "Amount", "Description"])
        
        incomes = Income.objects.filter(user=user)
        for inc in incomes:
            writer.writerow(["Income", inc.date.strftime('%Y-%m-%d'), inc.source, inc.category.name, inc.amount, inc.description or ""])
            
        expenses = Expense.objects.filter(user=user)
        for exp in expenses:
            writer.writerow(["Expense", exp.date.strftime('%Y-%m-%d'), exp.category.name, exp.category.name, exp.amount, exp.description or ""])
            
        return response
